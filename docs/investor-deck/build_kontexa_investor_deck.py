"""
Kontexa Investor Pitch Deck — story-first redesign (premium visual style preserved).
Story: Problem → Why it matters → Existing solutions → Why they fail → Kontexa →
How it works → Demo → Client transformation → Market → Advantage → Vision
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Kontexa_Investor_Pitch_Deck_v2.pptx"

BG = RGBColor(0x0B, 0x12, 0x20)
SURFACE = RGBColor(0x12, 0x1A, 0x2B)
SURFACE2 = RGBColor(0x18, 0x24, 0x3A)
CARD = RGBColor(0x15, 0x1F, 0x33)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
SOFT = RGBColor(0xCB, 0xD5, 0xE1)
RED = RGBColor(0xF4, 0x3F, 0x5E)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GREEN = RGBColor(0x10, 0xB9, 0x81)
LINE = RGBColor(0x2A, 0x3A, 0x55)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 12


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def rect(slide, left, top, w, h, fill, line=None, radius=0.1):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def oval(slide, left, top, size, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def run(r, text, size=18, bold=False, color=WHITE):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Segoe UI"


def text(slide, left, top, w, h, s, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run(p.add_run(), s, size=size, bold=bold, color=color)
    return box


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def footer(slide, page):
    text(slide, Inches(0.55), Inches(7.12), Inches(5), Inches(0.25),
         "Kontexa  ·  Confidential", size=10, color=MUTED)
    text(slide, Inches(11.5), Inches(7.12), Inches(1.3), Inches(0.25),
         f"{page} / {TOTAL}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def logo(slide):
    slide.shapes.add_picture(str(ASSETS / "kontexa-mark.png"), Inches(0.55), Inches(0.32), height=Inches(0.4))


def accent(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(0.85), Inches(0.07))
    sh.fill.solid()
    sh.fill.fore_color.rgb = TEAL
    sh.line.fill.background()


def title_block(slide, title, eyebrow=None):
    logo(slide)
    accent(slide)
    if eyebrow:
        text(slide, Inches(0.55), Inches(1.1), Inches(12), Inches(0.3), eyebrow, size=12, color=TEAL)
        text(slide, Inches(0.55), Inches(1.35), Inches(12), Inches(0.55), title, size=26, bold=True)
    else:
        text(slide, Inches(0.55), Inches(1.15), Inches(12), Inches(0.55), title, size=26, bold=True)


def arrow(slide, left, top, w=Inches(0.32), h=Inches(0.2)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = LINE
    sh.line.fill.background()
    return sh


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ========== 1 COVER ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), SLIDE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    s.shapes.add_picture(str(ASSETS / "kontexa-wordmark.png"), Inches(1.0), Inches(1.7), width=Inches(8.0))
    text(s, Inches(1.05), Inches(4.2), Inches(11), Inches(0.7),
         "Enterprise leaders ask “why.” Analysts write SQL. Decisions wait.",
         size=22, color=SOFT)
    text(s, Inches(1.05), Inches(4.95), Inches(11), Inches(0.5),
         "Kontexa turns business questions into warehouse-proven explanations.",
         size=16, color=MUTED)
    rect(s, Inches(1.05), Inches(5.7), Inches(5.0), Inches(0.85), SURFACE2)
    text(s, Inches(1.25), Inches(5.88), Inches(4.6), Inches(0.5),
         "Founder  ·  Parth Harduley", size=16, bold=True)
    text(s, Inches(7.0), Inches(5.95), Inches(5), Inches(0.4),
         "Investor & Enterprise Briefing", size=14, color=MUTED)
    notes(s, "Open with the business tension: executives need why-answers; analytics cycles delay decisions. "
             "Kontexa is warehouse-native decision intelligence — not another dashboard.")

    # ========== 2 PROBLEM WORKFLOW ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Every urgent “why” becomes another analytics cycle.",
                "The enterprise problem")
    text(s, Inches(0.55), Inches(2.05), Inches(12), Inches(0.4),
         "VP Sales · VP Marketing · CFO · COO · Strategy · PMs · Business Analysts",
         size=13, color=MUTED)

    steps = [
        ("1", "Business\nquestion", "“Why did revenue\ndecline?”"),
        ("2", "Ticket to\nanalytics", "Queue · priority\n· wait"),
        ("3", "SQL &\ndashboards", "Queries · reports\n· rebuilds"),
        ("4", "Follow-up\nquestions", "New tickets.\nSame loop."),
        ("5", "Delayed\ndecision", "Days or weeks\nlater"),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = Inches(0.4 + i * 2.55)
        color = RED if i == 4 else (AMBER if i >= 3 else CARD)
        rect(s, x, Inches(2.65), Inches(2.35), Inches(3.15), color if i == 4 else CARD, LINE)
        oval(s, x + Inches(0.9), Inches(2.85), Inches(0.5), TEAL if i < 4 else RED)
        text(s, x + Inches(0.9), Inches(2.92), Inches(0.5), Inches(0.4),
             n, size=14, bold=True, color=BG, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.15), Inches(3.55), Inches(2.05), Inches(0.7),
             title, size=14, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.15), Inches(4.4), Inches(2.05), Inches(0.9),
             body, size=12, color=SOFT if i == 4 else MUTED, align=PP_ALIGN.CENTER)
        if i < 4:
            arrow(s, x + Inches(2.2), Inches(4.0))
    text(s, Inches(0.55), Inches(6.15), Inches(12), Inches(0.4),
         "The bottleneck is not data. It is the human analytics loop between question and decision.",
         size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    footer(s, 2)
    notes(s, "Walk the workflow slowly. Emphasize personas asking real questions. "
             "Each new question restarts SQL/dashboard work. This is the cost center Kontexa collapses.")

    # ========== 3 WHY IT MATTERS ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Time-to-decision is a P&L variable.", "Why this matters")
    impacts = [
        ("Revenue", "Missed interventions while teams wait for the next report."),
        ("Profitability", "Margin leaks persist across another analytics sprint."),
        ("Retention", "Churn drivers stay opaque until after the damage."),
        ("Advantage", "Competitors who decide faster set the market pace."),
    ]
    for i, (t, b) in enumerate(impacts):
        x = Inches(0.55 + (i % 2) * 6.25)
        y = Inches(2.2 + (i // 2) * 2.0)
        rect(s, x, y, Inches(5.95), Inches(1.75), CARD, LINE)
        rect(s, x, y, Inches(0.12), Inches(1.75), TEAL if i % 2 == 0 else BLUE)
        text(s, x + Inches(0.4), y + Inches(0.35), Inches(5.2), Inches(0.4), t, size=20, bold=True)
        text(s, x + Inches(0.4), y + Inches(0.9), Inches(5.2), Inches(0.6), b, size=14, color=MUTED)
    footer(s, 3)
    notes(s, "Connect decision latency to CFO/COO language. This is not an IT convenience story — it is operating performance.")

    # ========== 4 EXISTING SOLUTIONS ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Enterprises already bought tools. The loop remains.", "Existing solutions")
    rows = [
        ("Dashboards & BI", "Monitor known KPIs well", "Cannot investigate new “why” questions alone"),
        ("Analytics teams", "Deep SQL expertise", "Finite capacity; every ask becomes a ticket"),
        ("Text-to-SQL chat", "Faster query drafting", "Brittle answers; weak evidence & governance"),
        ("Generic AI assistants", "Fluent explanations", "No warehouse proof; risk of invented metrics"),
    ]
    rect(s, Inches(0.55), Inches(2.1), Inches(12.2), Inches(0.5), SURFACE2)
    for i, h in enumerate(["Approach", "What it does well", "Where the loop breaks"]):
        text(s, Inches(0.7 + i * 4.0), Inches(2.2), Inches(3.8), Inches(0.35), h, size=12, bold=True, color=TEAL)
    for r, (a, b, c) in enumerate(rows):
        y = Inches(2.7 + r * 0.9)
        rect(s, Inches(0.55), y, Inches(12.2), Inches(0.8), CARD if r % 2 == 0 else SURFACE)
        text(s, Inches(0.7), y + Inches(0.22), Inches(3.8), Inches(0.4), a, size=14, bold=True)
        text(s, Inches(4.7), y + Inches(0.22), Inches(3.8), Inches(0.4), b, size=13, color=MUTED)
        text(s, Inches(8.7), y + Inches(0.22), Inches(3.8), Inches(0.4), c, size=13, color=SOFT)
    footer(s, 4)
    notes(s, "Do not attack customers for buying Tableau/Power BI. Acknowledge value, then show the structural gap for investigation.")

    # ========== 5 WHY DASHBOARDS AREN'T ENOUGH ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Why dashboards are no longer enough", "The monitoring vs reasoning gap")
    rect(s, Inches(0.55), Inches(2.15), Inches(5.9), Inches(4.3), CARD, LINE)
    text(s, Inches(0.85), Inches(2.4), Inches(5.3), Inches(0.4), "Dashboards excel at…", size=16, bold=True, color=GREEN)
    for i, item in enumerate([
        "Tracking known KPIs in real time",
        "Shared visibility across teams",
        "Trend monitoring & alerting",
        "Operational scorecards",
    ]):
        text(s, Inches(0.85), Inches(3.05 + i * 0.7), Inches(5.3), Inches(0.5), "●  " + item, size=15, color=SOFT)

    rect(s, Inches(6.85), Inches(2.15), Inches(5.9), Inches(4.3), CARD, LINE)
    text(s, Inches(7.15), Inches(2.4), Inches(5.3), Inches(0.4), "They struggle with…", size=16, bold=True, color=AMBER)
    for i, item in enumerate([
        "New “why / explain / investigate” asks",
        "Competing hypotheses under test",
        "Evidence tied to warehouse rows",
        "Board-ready causal honesty",
    ]):
        text(s, Inches(7.15), Inches(3.05 + i * 0.7), Inches(5.3), Inches(0.5), "●  " + item, size=15, color=SOFT)

    text(s, Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.35),
         "Kontexa is the reasoning layer that complements dashboards — and grounds AI assistants in warehouse truth.",
         size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    footer(s, 5)
    notes(s, "Position carefully: dashboards are not obsolete. They monitor. Kontexa reasons. Together they form the modern analytics stack.")

    # ========== 6 WHY NOT CHATGPT ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Why not ChatGPT / generic AI?", "Enterprise decision intelligence requires more than fluency")
    headers = ["Requirement", "Generic AI", "Kontexa"]
    matrix = [
        ["Warehouse connectivity", "Usually absent / ad hoc", "Native evidence execution"],
        ["Governed business catalogue", "Prompt memory ≠ semantics", "Approved catalogue authority"],
        ["Evidence-backed reasoning", "Narrative can invent facts", "Claims tied to query results"],
        ["Explainable conclusions", "Hard to audit", "Hypotheses · purpose · unknowns"],
        ["Decision-grade confidence", "Sounds certain", "Honest insufficiency when needed"],
    ]
    rect(s, Inches(0.55), Inches(2.1), Inches(12.2), Inches(0.5), SURFACE2)
    xs = [0.7, 4.5, 8.6]
    for i, h in enumerate(headers):
        text(s, Inches(xs[i]), Inches(2.2), Inches(3.5), Inches(0.35), h, size=13, bold=True, color=TEAL)
    for r, row in enumerate(matrix):
        y = Inches(2.7 + r * 0.7)
        rect(s, Inches(0.55), y, Inches(12.2), Inches(0.62), CARD if r % 2 == 0 else SURFACE)
        for i, cell in enumerate(row):
            color = GREEN if i == 2 else (MUTED if i == 1 else WHITE)
            text(s, Inches(xs[i]), y + Inches(0.15), Inches(3.5), Inches(0.4), cell, size=13,
                 bold=(i == 0 or i == 2), color=color)
    footer(s, 6)
    notes(s, "Investors will ask this. Be crisp: LLMs are the interface; enterprises need a governed reasoning engine connected to their warehouse.")

    # ========== 7 KONTEXA ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Kontexa: from business question to proven explanation.", "The solution")
    text(s, Inches(0.55), Inches(2.05), Inches(12), Inches(0.4),
         "An evidence-driven, warehouse-native decision intelligence platform.",
         size=16, color=SOFT)
    pillars = [
        ("Ask", "Leaders ask in business language — no ticket required."),
        ("Plan", "GPT builds an investigation with competing hypotheses."),
        ("Prove", "Warehouse evidence validates or weakens each claim."),
        ("Decide", "Executives get explanation, confidence, and unknowns."),
    ]
    for i, (t, b) in enumerate(pillars):
        x = Inches(0.55 + i * 3.15)
        rect(s, x, Inches(2.7), Inches(2.95), Inches(3.4), CARD, LINE)
        oval(s, x + Inches(1.1), Inches(3.0), Inches(0.6), TEAL if i % 2 == 0 else BLUE)
        text(s, x + Inches(1.1), Inches(3.1), Inches(0.6), Inches(0.4),
             str(i + 1), size=16, bold=True, color=BG, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.2), Inches(3.9), Inches(2.55), Inches(0.45), t, size=18, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.2), Inches(4.5), Inches(2.55), Inches(1.2), b, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 7)
    notes(s, "One-liner: Kontexa collapses the analytics ticket loop into an evidence program that executives can trust.")

    # ========== 8 HOW IT WORKS ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "How Kontexa works", "Six steps from question to executive explanation")
    steps = [
        ("1", "Business\nquestion", "Natural language\nfrom the leader"),
        ("2", "Catalogue\nunderstanding", "Approved semantics\n& definitions"),
        ("3", "Investigation\nplanning", "Competing\nhypotheses"),
        ("4", "Warehouse\nevidence", "SQL evidence\nrequests executed"),
        ("5", "Evidence\nvalidation", "Purpose-tagged\nresults packed"),
        ("6", "Executive\nexplanation", "Confidence &\nunknowns stated"),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = Inches(0.35 + i * 2.15)
        rect(s, x, Inches(2.3), Inches(2.0), Inches(3.7), CARD, LINE)
        oval(s, x + Inches(0.7), Inches(2.55), Inches(0.5), TEAL)
        text(s, x + Inches(0.7), Inches(2.62), Inches(0.5), Inches(0.4),
             n, size=14, bold=True, color=BG, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.1), Inches(3.3), Inches(1.8), Inches(0.8),
             title, size=13, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.1), Inches(4.3), Inches(1.8), Inches(1.2),
             body, size=11, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            arrow(s, x + Inches(1.85), Inches(3.9), Inches(0.28), Inches(0.18))
    text(s, Inches(0.55), Inches(6.3), Inches(12), Inches(0.35),
         "GPT plans. The warehouse proves. Narrative never invents support.",
         size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    footer(s, 8)
    notes(s, "This is the product architecture told as a business process. Emphasize hypotheses and insufficiency as trust features.")

    # ========== 9 PRODUCT DEMO ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Product experience", "From ask to board-ready insight")
    mocks = [
        ("mock-home.png", "Home"),
        ("mock-ask.png", "Ask WHY"),
        ("mock-answer.png", "Evidence answer"),
        ("mock-charts.png", "Charts"),
        ("mock-investigation.png", "Investigation"),
        ("mock-executive.png", "Executive insight"),
    ]
    for i, (fn, label) in enumerate(mocks):
        x = Inches(0.4 + (i % 3) * 4.25)
        y = Inches(2.05 + (i // 3) * 2.4)
        path = ASSETS / fn
        rect(s, x, y, Inches(4.05), Inches(2.05), SURFACE, LINE, radius=0.06)
        if path.exists():
            s.shapes.add_picture(str(path), x + Inches(0.1), y + Inches(0.1), width=Inches(3.85))
        text(s, x, y + Inches(2.1), Inches(4.05), Inches(0.28),
             label, size=12, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 9)
    notes(s, "If presenting live, swap mockups for product screens. Story: ask → plan → evidence → charts → hypotheses → executive card.")

    # ========== 10 WHAT CHANGES FOR THE CLIENT ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "What changes for the client", "The transformation that matters")

    # BEFORE lane
    rect(s, Inches(0.4), Inches(2.05), Inches(12.5), Inches(2.15), CARD, RED)
    text(s, Inches(0.6), Inches(2.2), Inches(2), Inches(0.35), "BEFORE", size=14, bold=True, color=RED)
    before = [
        ("Executive", "asks why"),
        ("Analyst\nticket", "queued"),
        ("SQL &\nreports", "built"),
        ("Dashboard", "shared"),
        ("Follow-ups", "restart"),
        ("Decision", "delayed"),
    ]
    for i, (t, b) in enumerate(before):
        x = Inches(0.55 + i * 2.05)
        rect(s, x, Inches(2.65), Inches(1.9), Inches(1.3), SURFACE, LINE)
        text(s, x + Inches(0.05), Inches(2.8), Inches(1.8), Inches(0.55), t, size=12, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.05), Inches(3.4), Inches(1.8), Inches(0.35), b, size=11, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            arrow(s, x + Inches(1.75), Inches(3.15), Inches(0.25), Inches(0.16))

    # AFTER lane
    rect(s, Inches(0.4), Inches(4.45), Inches(12.5), Inches(2.15), CARD, GREEN)
    text(s, Inches(0.6), Inches(4.6), Inches(3), Inches(0.35), "AFTER KONTEXA", size=14, bold=True, color=GREEN)
    after = [
        ("Executive", "asks why"),
        ("Kontexa", "plans"),
        ("Warehouse\ninvestigation", "runs"),
        ("Evidence\nexplanation", "proven"),
        ("Confident\ndecision", "minutes"),
    ]
    for i, (t, b) in enumerate(after):
        x = Inches(0.7 + i * 2.45)
        rect(s, x, Inches(5.05), Inches(2.25), Inches(1.3), SURFACE, TEAL)
        text(s, x + Inches(0.05), Inches(5.2), Inches(2.15), Inches(0.55), t, size=12, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.05), Inches(5.8), Inches(2.15), Inches(0.35), b, size=11, color=SOFT, align=PP_ALIGN.CENTER)
        if i < 4:
            arrow(s, x + Inches(2.1), Inches(5.55), Inches(0.28), Inches(0.16))
    footer(s, 10)
    notes(s, "Strongest visual slide. Pause here. Before = ticket factory. After = minutes to evidence-backed decision.")

    # ========== 11 MARKET ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    title_block(s, "Market opportunity", "TAM · SAM · SOM")
    markets = [
        (Inches(2.0), Inches(2.15), Inches(9.3), Inches(1.3), "TAM", "$40.1B",
         "Global BI Software (2025) → $81.5B by 2033 · CAGR 9.3%  ·  Grand View Research", TEAL),
        (Inches(2.8), Inches(3.65), Inches(7.7), Inches(1.3), "SAM", "$19.4B",
         "Decision Intelligence (2025) → $57.8B by 2032 · CAGR 16.9%  ·  Fortune Business Insights", BLUE),
        (Inches(3.8), Inches(5.15), Inches(5.7), Inches(1.25), "SOM", "$0.6–1.0B",
         "Beachhead: large-enterprise warehouse-native decision layer (5–7 yr)", GREEN),
    ]
    for left, top, w, h, label, value, desc, color in markets:
        rect(s, left, top, w, h, CARD, color)
        text(s, left + Inches(0.25), top + Inches(0.2), Inches(1.2), Inches(0.3), label, size=12, bold=True, color=color)
        text(s, left + Inches(1.5), top + Inches(0.15), Inches(2.2), Inches(0.45), value, size=24, bold=True)
        text(s, left + Inches(3.7), top + Inches(0.3), w - Inches(4.0), Inches(0.7), desc, size=12, color=MUTED)
    footer(s, 11)
    notes(s,
          "Sources: Grand View Research BI Software Market (2025–2033); Fortune Business Insights Decision Intelligence (2025–2032); "
          "FMI Augmented Analytics adjacent; Gartner software spend ~$1.25T (2025). SOM is Kontexa bottom-up beachhead estimate.")

    # ========== 12 VISION ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.16), SLIDE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    s.shapes.add_picture(str(ASSETS / "kontexa-mark.png"), Inches(1.0), Inches(0.9), height=Inches(0.55))
    text(s, Inches(1.75), Inches(1.0), Inches(8), Inches(0.45), "Kontexa", size=26, bold=True)
    text(s, Inches(1.0), Inches(1.8), Inches(11), Inches(0.55),
         "The warehouse-native reasoning engine behind enterprise decisions.",
         size=24, bold=True)
    text(s, Inches(1.0), Inches(2.55), Inches(11), Inches(0.7),
         "As leaders interact through AI assistants (Claude and others), Kontexa becomes the governed "
         "evidence layer — not another standalone dashboard.",
         size=15, color=SOFT)

    cards = [
        ("Near term", "Own WHY / EXPLAIN for cloud warehouse enterprises."),
        ("Medium term", "Embed decision workflows in finance, growth, and ops."),
        ("Long term", "Power assistant experiences with warehouse-proven reasoning."),
    ]
    for i, (t, b) in enumerate(cards):
        x = Inches(1.0 + i * 3.9)
        rect(s, x, Inches(3.6), Inches(3.6), Inches(1.7), CARD, LINE)
        text(s, x + Inches(0.25), Inches(3.85), Inches(3.1), Inches(0.35), t, size=13, bold=True, color=TEAL)
        text(s, x + Inches(0.25), Inches(4.35), Inches(3.1), Inches(0.7), b, size=13, color=SOFT)

    text(s, Inches(1.0), Inches(5.7), Inches(11), Inches(0.4),
         "Business Intelligence Made Easy", size=18, bold=True)
    text(s, Inches(1.0), Inches(6.2), Inches(11), Inches(0.35),
         "Parth Harduley  ·  Founder  ·  Let’s collapse the distance between question and decision.",
         size=14, color=MUTED)
    footer(s, 12)
    notes(s, "Close on platform vision: assistants as interface, Kontexa as reasoning engine. Ask for pilot / partnership next step.")

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
